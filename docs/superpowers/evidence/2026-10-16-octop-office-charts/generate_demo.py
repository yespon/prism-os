from pathlib import Path
import hashlib,io,json
from openpyxl import load_workbook
from pptx import Presentation
from deerflow.tools.builtins.office_xlsx import create_xlsx,edit_xlsx,inspect_xlsx
from deerflow.tools.builtins.office_pptx import create_pptx,edit_pptx,inspect_pptx
HERE=Path(__file__).resolve().parent
rows=[['季度','预算','实际'],['第一季',120,110],['第二季',160,145],['第三季',180,175],['第四季',200,190]]
charts=[{'kind':kind,'title':title,'categories':'A2:A5','values':'B1:B5' if kind=='pie' else 'B1:C5','anchor':anchor} for kind,title,anchor in [('bar','预算与实际','E2'),('line','季度走势','E18'),('pie','预算分布','E34')]]
book=create_xlsx({'sheets':[{'name':'季度预算','header':True,'filter':True,'rows':rows,'charts':charts}]})
book=edit_xlsx(book,[{'sheet':'季度预算','cell':'B2','value':125}])
slides=[]
for kind in ['bar','line','pie']:
    series=[{'name':'预算','values':[125,160,180,200]}]
    if kind!='pie':series.append({'name':'实际','values':[110,145,175,190]})
    slides.append({'title':{'bar':'预算与实际','line':'季度走势','pie':'预算分布'}[kind],'chart':{'kind':kind,'title':'本地可编辑图表','categories':[r[0] for r in rows[1:]],'series':series},'notes':'实际文件示例；桌面渲染仍待验收，未调用模型或外部工具。'})
deck=create_pptx({'slides':slides})
deck=edit_pptx(deck,{'预算与实际':'项目预算与实际'})
assert len(load_workbook(io.BytesIO(book)).active._charts)==3
assert len(Presentation(io.BytesIO(deck)).slides)==3
proof={'scope':'Actual bounded independent Office chart packages, not desktop rendering or model behavior acceptance','files':{}}
for name,data,inspect in [('demo-charts.xlsx',book,inspect_xlsx),('demo-charts.pptx',deck,inspect_pptx)]:
    path=HERE/name
    assert not path.exists()
    path.write_bytes(data)
    proof['files'][name]={'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data),'inspected':inspect(data)}
(HERE/'demo.json').write_text(json.dumps(proof,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({name:{k:v for k,v in value.items() if k!='inspected'} for name,value in proof['files'].items()}))
