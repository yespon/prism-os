from pathlib import Path
import hashlib,io,json
from pptx import Presentation
from openpyxl import load_workbook
from deerflow.tools.builtins.office_pptx import create_pptx,edit_pptx_charts,inspect_pptx
HERE=Path(__file__).resolve().parent
spec={'slides':[{'title':'季度预算与实际','chart':{'kind':'bar','title':'原始预算','categories':['第一季','第二季'],'series':[{'name':'预算','values':[120,160]}]},'notes':'独立图表数据编辑演示'},{'title':'保留的其他内容','paragraphs':['编辑指定图表，保留其他页面','桌面渲染与高级组合图仍待验收']}]}
original=create_pptx(spec)
new={'kind':'bar','title':'更新后的预算与实际','categories':['第一季','第二季','第三季'],'series':[{'name':'预算','values':[125,165,190]},{'name':'实际','values':[110,150,180]}]}
changed=edit_pptx_charts(original,[{'slide':1,'chart':1,'document':new}])
actual=Presentation(io.BytesIO(changed))
chart=actual.slides[0].shapes[1].chart
assert list(chart.series[0].values)==[125,165,190]
book=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
assert book.active['C4'].value==180
assert inspect_pptx(original)['slides'][1]==inspect_pptx(changed)['slides'][1]
proof={'scope':'Actual OOXML charts/caches/editable workbook; native authority/download separately tested; no desktop rendering/model call','files':{}}
for name,data in [('demo-original.pptx',original),('demo-edited.pptx',changed)]:
 path=HERE/name
 assert not path.exists()
 path.write_bytes(data)
 proof['files'][name]={'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)}
proof['inspection']=inspect_pptx(changed)
(HERE/'demo.json').write_text(json.dumps(proof,indent=2,ensure_ascii=False)+'\n')
print(json.dumps(proof['files']))
