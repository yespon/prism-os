from pathlib import Path
import hashlib,io,json,sys
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[3]
sys.path.insert(0,str(ROOT/'backend'))
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from deerflow.tools.builtins.office_pptx import create_pptx,inspect_pptx
from deerflow.tools.builtins.office_pptx_images import add_pptx_images
canvas=Image.new('RGB',(720,300),'white');draw=ImageDraw.Draw(canvas)
for index,(label,value) in enumerate([('Q1',80),('Q2',140),('Q3',210),('Q4',170)]):
 x=90+index*150;draw.rectangle((x,260-value,x+80,260),fill=(12,80,180));draw.text((x+25,275),label,fill='black')
stream=io.BytesIO();canvas.save(stream,format='PNG');image=stream.getvalue()
source=create_pptx({'slides':[{'title':'Quarter report','paragraphs':['Authorized PNG picture','Aspect ratio and geometry preserved'],'notes':'Image inserted through the independent bounded module.'},{'title':'Editable chart','chart':{'kind':'bar','categories':['Q1','Q2'],'series':[{'name':'Count','values':[80,140]}]}}]})
result=add_pptx_images(source,[{'slide':1,'data':image,'left_inches':6,'top_inches':2,'width_inches':6,'alt':'Quarterly count chart'}])
assert inspect_pptx(source)==inspect_pptx(result)
actual=Presentation(io.BytesIO(result));assert any(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in actual.slides[0].shapes)
for name,data in [('demo-source.pptx',source),('demo-pictured.pptx',result),('demo-source.png',image)]:
 (HERE/name).write_bytes(data)
(HERE/'demo.json').write_text(json.dumps({'files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in HERE.glob('demo-*')},'actual_picture_verified':True,'desktop_rendered':False},indent=2)+'\n')
print('Actual PPTX image, chart and notes demo generated; no desktop rendering claim')
