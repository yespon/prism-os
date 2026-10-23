"""Actual PPTX pictures, unaffected parts and current native file authority."""

import importlib
import io
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from test_file_access_boundaries import owned_runtime as owned_fixture
from test_file_access_boundaries import selected_runtime  # noqa: F401
from test_file_access_http import file_client as client_fixture
from test_office_docx import withdraw
from test_skill_member_tools import configured, scoped_client  # noqa: F401
from test_skill_reference_transactions import durable_sqlite  # noqa: F401

from deerflow.tools.builtins.office_pptx import create_pptx, inspect_pptx, office_pptx_tool
from deerflow.tools.builtins.office_pptx_images import add_pptx_images

owned_runtime = owned_fixture
file_client = client_fixture
MODULE = importlib.import_module('deerflow.tools.builtins.office_pptx_images')
DOCUMENT = {'slides': [{'title': '图文报告', 'paragraphs': ['保留正文'], 'notes': '保留备注'}, {'title': '原图表', 'chart': {'kind': 'bar', 'categories': ['类别'], 'series': [{'name': '数值', 'values': [2]}]}}]}


def picture():
    stream = io.BytesIO()
    Image.new('RGB', (120, 60), (12, 80, 180)).save(stream, format='PNG')
    return stream.getvalue()


def test_real_picture_relationship_pixels_geometry_and_other_chart_unchanged():
    source = create_pptx(DOCUMENT)
    result = add_pptx_images(source, [{'slide': 1, 'data': picture(), 'width_inches': 3, 'left_inches': 1, 'top_inches': 2, 'alt': '实际图像'}])
    assert inspect_pptx(result) == inspect_pptx(source)
    actual = Presentation(io.BytesIO(result))
    image = next(shape for shape in actual.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert [image.left, image.top, image.width, image.height] == [914400, 1828800, 2743200, 1371600]
    assert image._element.nvPicPr.cNvPr.get('descr') == '实际图像'
    with Image.open(io.BytesIO(image.image.blob)) as pixels:
        assert pixels.size == (120, 60) and pixels.getpixel((0, 0)) == (12, 80, 180)
    with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(io.BytesIO(result)) as new:
        for name in old.namelist():
            if name.startswith(('ppt/charts/', 'ppt/embeddings/', 'ppt/notesSlides/')) or name.startswith('ppt/slides/slide2'):
                assert old.read(name) == new.read(name), name


def test_repeated_insertions_deduplicate_media_keep_distinct_shape_ids():
    source = create_pptx({'slides': [{'title': 'one'}]})
    first = add_pptx_images(source, [{'slide': 1, 'data': picture()}])
    second = add_pptx_images(first, [{'slide': 1, 'data': picture()}])
    actual = Presentation(io.BytesIO(second))
    ids = [s.shape_id for s in actual.slides[0].shapes]
    assert len(ids) == len(set(ids))
    assert sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in actual.slides[0].shapes) == 2
    with zipfile.ZipFile(io.BytesIO(second)) as package:
        assert len([n for n in package.namelist() if n.startswith('ppt/media/')]) == 1


@pytest.mark.parametrize('changes', [
    {'slide': 0}, {'slide': True}, {'slide': 3}, {'width_inches': float('nan')}, {'width_inches': True}, {'left_inches': -1}, {'top_inches': 7.5}, {'width_inches': 99}, {'alt': 'a' * 2001}, {'data': b'<svg/>'}, {'unknown': 1}
])
def test_invalid_images_or_off_slide_geometry_fail_explicitly(changes):
    item = {'slide': 1, 'data': picture(), **changes}
    with pytest.raises((ValueError, OSError)):
        add_pptx_images(create_pptx(DOCUMENT), [item])


def test_bounded_total_picture_count_across_repeated_edits():
    data = create_pptx({'slides': [{'title': 'bounded'}]})
    for _ in range(4):
        data = add_pptx_images(data, [{'slide': 1, 'data': picture()}] * 16)
    with pytest.raises(ValueError, match='picture limit'):
        add_pptx_images(data, [{'slide': 1, 'data': picture()}])


def test_native_insert_download_foreign_image_and_actor_denial(file_client):
    client, _, _, runtime = file_client
    output = Path(runtime.state['thread_data']['outputs_path'])
    uploads = Path(runtime.state['thread_data']['uploads_path'])
    (uploads / 'image.png').write_bytes(picture())
    (output / 'source.pptx').write_bytes(create_pptx(DOCUMENT))
    result = office_pptx_tool.func(runtime, 'add_images', source_path='/mnt/user-data/outputs/source.pptx', output_path='/mnt/user-data/outputs/images.pptx',
        images=[{'slide': 1, 'source_path': '/mnt/user-data/uploads/image.png'}])
    assert result['status'] == 'ok'
    response = client.get('/api/threads/member-thread/artifacts/mnt/user-data/outputs/images.pptx?download=true')
    assert response.status_code == 200
    assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in Presentation(io.BytesIO(response.content)).slides[0].shapes)
    assert client.get('/api/threads/member-thread/artifacts/mnt/user-data/outputs/images.pptx', headers={'x-user': 'user-b'}).status_code == 404
    private = output.parent / 'private.png'
    private.write_bytes(picture())
    assert office_pptx_tool.func(runtime, 'add_images', source_path='/mnt/user-data/outputs/source.pptx', output_path='/mnt/user-data/outputs/foreign.pptx', images=[{'slide': 1, 'source_path': str(private)}])['status'] == 'unavailable'
    assert not (output / 'foreign.pptx').exists()
    assert office_pptx_tool.func(runtime, 'add_images', source_path='/mnt/user-data/outputs/source.pptx', output_path='/mnt/user-data/outputs/images.pptx',
        images=[{'slide': 1, 'source_path': '/mnt/user-data/uploads/image.png'}])['status'] == 'unavailable'
    assert (output / 'images.pptx').read_bytes() == response.content


@pytest.mark.parametrize('point', ['before', 'during'])
def test_current_withdrawal_discards_presentation_before_publication(owned_runtime, monkeypatch, point):
    factory, _, _, _, runtime = owned_runtime
    output = Path(runtime.state['thread_data']['outputs_path'])
    uploads = Path(runtime.state['thread_data']['uploads_path'])
    (output / 'source.pptx').write_bytes(create_pptx(DOCUMENT))
    (uploads / 'image.png').write_bytes(picture())
    if point == 'before':
        withdraw(factory)
    else:
        original = MODULE.add_pptx_images
        def revoke(*args):
            result = original(*args)
            withdraw(factory)
            return result
        monkeypatch.setattr(MODULE, 'add_pptx_images', revoke)
    result = office_pptx_tool.func(runtime, 'add_images', source_path='/mnt/user-data/outputs/source.pptx', output_path='/mnt/user-data/outputs/revoked.pptx', images=[{'slide': 1, 'source_path': '/mnt/user-data/uploads/image.png'}])
    assert result == {'status': 'unavailable', 'inspection_required': True, 'retry': False}
    assert not (output / 'revoked.pptx').exists()


def test_image_schema_does_not_expose_runtime():
    schema = office_pptx_tool.tool_call_schema.model_json_schema()['properties']
    assert 'images' in schema and 'runtime' not in schema
