"""Real chart packages, editable data and original current-resource output gates."""

import copy
import importlib
import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from test_file_access_boundaries import owned_runtime as owned_fixture
from test_file_access_boundaries import selected_runtime  # noqa: F401
from test_file_access_http import file_client as client_fixture
from test_office_docx import withdraw
from test_skill_member_tools import configured, scoped_client  # noqa: F401
from test_skill_reference_transactions import durable_sqlite  # noqa: F401

from deerflow.tools.builtins.office_pptx import create_pptx, edit_pptx, inspect_pptx, office_pptx_tool
from deerflow.tools.builtins.office_xlsx import create_xlsx, edit_xlsx, inspect_xlsx, office_xlsx_tool

owned_runtime = owned_fixture
file_client = client_fixture
BOOK = {
    "sheets": [{"name": "季度预算", "header": True, "rows": [["季度", "预算", "实际"], ["第一季", 12.5, 10], ["第二季", 20, 17]], "charts": [{"kind": "bar", "title": "预算与实际", "categories": "A2:A3", "values": "B1:C3", "anchor": "E2"}]}]
}
DECK = {"slides": [{"title": "预算报告", "notes": "数据说明", "chart": {"kind": "bar", "title": "预算与实际", "categories": ["第一季", "第二季"], "series": [{"name": "预算", "values": [12.5, 20]}, {"name": "实际", "values": [10, 17]}]}}]}


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_spreadsheet_charts_reference_local_cells_and_survive_edit(kind):
    document = copy.deepcopy(BOOK)
    spec = document["sheets"][0]["charts"][0]
    spec["kind"] = kind
    if kind == "pie":
        spec["values"] = "B1:B3"
    original = create_xlsx(document)
    actual = load_workbook(io.BytesIO(original))
    chart = actual.active._charts[0]
    assert chart.series[0].val.numRef.f == "'季度预算'!$B$2:$B$3"
    assert chart.series[0].cat.numRef.f == "'季度预算'!$A$2:$A$3"
    with zipfile.ZipFile(io.BytesIO(original)) as package:
        xml = ET.fromstring(package.read("xl/charts/chart1.xml"))
        assert any(node.text == "预算与实际" for node in xml.iter())
    changed = edit_xlsx(original, [{"sheet": "季度预算", "cell": "B2", "value": 42}])
    edited = load_workbook(io.BytesIO(changed))
    assert edited.active["B2"].value == 42 and len(edited.active._charts) == 1
    assert edited.active._charts[0].series[0].val.numRef.f == chart.series[0].val.numRef.f
    assert actual.active["B2"].value == 12.5
    assert inspect_xlsx(changed)["sheets"][0]["charts"][0]["series_count"] == (1 if kind == "pie" else 2)


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_presentation_charts_have_actual_cached_values_and_editable_workbook(kind):
    document = copy.deepcopy(DECK)
    spec = document["slides"][0]["chart"]
    spec["kind"] = kind
    if kind == "pie":
        spec["series"] = spec["series"][:1]
    original = create_pptx(document)
    deck = Presentation(io.BytesIO(original))
    chart = deck.slides[0].shapes[1].chart
    assert chart.series[0].values == (12.5, 20)
    assert [item.label for item in chart.plots[0].categories] == ["第一季", "第二季"]
    workbook = load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
    assert workbook.active["A2"].value == "第一季" and workbook.active["B2"].value == 12.5
    changed = edit_pptx(original, {"预算报告": "年度报告", "数据说明": "保留图表"})
    edited = Presentation(io.BytesIO(changed))
    assert edited.slides[0].shapes[0].text == "年度报告"
    assert edited.slides[0].shapes[1].chart.part.chart_workbook.xlsx_part.blob == chart.part.chart_workbook.xlsx_part.blob
    result = inspect_pptx(changed)
    assert result["slides"][0]["charts"][0]["series"][0]["values"] == [12.5, 20]
    assert result["rendered"] is False


@pytest.mark.parametrize("mutation", ["external", "formula", "mismatch", "nonfinite", "pie-series", "anchor", "count"])
def test_chart_references_and_data_are_bounded_and_local(mutation):
    book, deck = copy.deepcopy(BOOK), copy.deepcopy(DECK)
    x, p = book["sheets"][0]["charts"][0], deck["slides"][0]["chart"]
    if mutation == "external":
        x["values"] = "'[foreign.xlsx]Sheet'!B1:C3"
        p["url"] = "https://invalid"
    elif mutation == "formula":
        book["sheets"][0]["rows"][1][1] = {"formula": "=SUM(B3)"}
        p["series"][0]["values"][0] = "=SUM(B3)"
    elif mutation == "mismatch":
        x["values"] = "B2:C3"
        p["series"][0]["values"].pop()
    elif mutation == "nonfinite":
        book["sheets"][0]["rows"][1][1] = float("nan")
        p["series"][0]["values"][0] = float("inf")
    elif mutation == "pie-series":
        x["kind"] = p["kind"] = "pie"
    elif mutation == "anchor":
        x["anchor"] = "ZZZ10001"
        p["categories"] = ["too many"] * 201
    else:
        book["sheets"][0]["charts"] *= 17
        deck["slides"] *= 33
    for create, document in [(create_xlsx, book), (create_pptx, deck)]:
        with pytest.raises(ValueError):
            create(document)


def test_presentation_nested_chart_workbook_rejects_external_relationships():
    original = create_pptx(DECK)
    with zipfile.ZipFile(io.BytesIO(original)) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    path = next(name for name in parts if name.startswith("ppt/embeddings/"))
    with zipfile.ZipFile(io.BytesIO(parts[path])) as workbook:
        cells = {name: workbook.read(name) for name in workbook.namelist()}
    cells["xl/_rels/workbook.xml.rels"] = b'<Relationships><Relationship TargetMode="External" Target="file:///private"/></Relationships>'
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as package:
        for name, data in cells.items():
            package.writestr(name, data)
    parts[path] = inner.getvalue()
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    with pytest.raises(ValueError, match="External Office"):
        inspect_pptx(outer.getvalue())


def test_real_chart_outputs_use_original_download_and_cross_member_gates(file_client):
    client, _, _, runtime = file_client
    for suffix, tool, document in [("xlsx", office_xlsx_tool, BOOK), ("pptx", office_pptx_tool, DECK)]:
        path = "/mnt/user-data/outputs/chart." + suffix
        assert tool.func(runtime, "create", output_path=path, document=document)["status"] == "ok"
        response = client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/chart." + suffix + "?download=true")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        assert b"chart" in response.content
        assert client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/chart." + suffix, headers={"x-user": "user-b"}).status_code == 404
        assert tool.func(runtime, "create", output_path=path, document=document)["status"] == "unavailable"


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_chart_generation_late_withdrawal_does_not_publish(owned_runtime, monkeypatch, kind):
    factory, _, _, _, runtime = owned_runtime
    module = importlib.import_module("deerflow.tools.builtins.office_" + kind)
    original = getattr(module, "create_" + kind)

    def revoke(document):
        result = original(document)
        withdraw(factory)
        return result

    monkeypatch.setattr(module, "create_" + kind, revoke)
    tool, document = (office_xlsx_tool, BOOK) if kind == "xlsx" else (office_pptx_tool, DECK)
    result = tool.func(runtime, "create", output_path="/mnt/user-data/outputs/chart." + kind, document=document)
    assert result == {"status": "unavailable", "inspection_required": True, "retry": False}
    assert not (Path(runtime.state["thread_data"]["outputs_path"]) / ("chart." + kind)).exists()


def test_presentation_chart_labels_remain_literal_in_embedded_editable_cells():
    spec = copy.deepcopy(DECK)
    spec["slides"][0]["chart"]["categories"][0] = '=WEBSERVICE("https://invalid")'
    spec["slides"][0]["chart"]["series"][0]["name"] = "https://invalid/label"
    deck = Presentation(io.BytesIO(create_pptx(spec)))
    blob = deck.slides[0].shapes[1].chart.part.chart_workbook.xlsx_part.blob
    book = load_workbook(io.BytesIO(blob))
    assert book.active["A2"].data_type == "s"
    assert book.active["A2"].value == '=WEBSERVICE("https://invalid")'
    assert book.active["B1"].hyperlink is None
