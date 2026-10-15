"""Actual presentation and PDF engines plus original owned output authorization."""

import asyncio
import io
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from sqlalchemy import select
from test_file_access_boundaries import owned_runtime as owned_fixture
from test_file_access_boundaries import selected_runtime  # noqa: F401
from test_file_access_http import file_client as client_fixture
from test_skill_member_tools import configured, scoped_client  # noqa: F401
from test_skill_reference_transactions import durable_sqlite  # noqa: F401

from deerflow.database.models import TenantMembership
from deerflow.tools.builtins.office_pdf import create_pdf, inspect_pdf, merge_pdf, office_pdf_tool, render_pdf_page, select_pdf_pages
from deerflow.tools.builtins.office_pptx import create_pptx, edit_pptx, inspect_pptx, office_pptx_tool

owned_runtime = owned_fixture
file_client = client_fixture
SLIDES = {"slides": [{"title": "季度总结", "paragraphs": ["已完成", "下一步"], "notes": "汇报备注"}, {"title": "数据", "table": [["部门", "预算"], ["研发", "12"]]}]}
PDF = {"pages": [["PrismOS report", "Budget: 12.50", "Cafe: caf\u00e9"], ["Second page", "Actual independent PDF"]]}


def test_actual_presentation_text_tables_notes_and_existing_style_roundtrip():
    source = create_pptx(SLIDES)
    with zipfile.ZipFile(io.BytesIO(source)) as package:
        assert not any(name.endswith(".bin") for name in package.namelist())
    actual = Presentation(io.BytesIO(source))
    assert len(actual.slides) == 2
    assert actual.slides[0].shapes[0].text == "季度总结"
    assert actual.slides[0].shapes[0].text_frame.paragraphs[0].font.bold
    assert actual.slides[1].shapes[1].table.cell(1, 0).text == "研发"
    changed = edit_pptx(source, {"季度总结": "年度总结", "研发": "平台", "汇报备注": "演示备注"})
    result = inspect_pptx(changed)
    assert result["slides"][0]["texts"][0] == "年度总结"
    assert result["slides"][0]["notes"] == "演示备注"
    assert result["slides"][1]["tables"][0][1][0] == "平台"
    assert inspect_pptx(source)["slides"][0]["texts"][0] == "季度总结"
    assert Presentation(io.BytesIO(changed)).slides[0].shapes[0].text_frame.paragraphs[0].font.bold


def test_presentation_growth_and_shape_limits_are_explicit():
    source = create_pptx({"slides": [{"paragraphs": ["a" * 100]}]})
    with pytest.raises(ValueError, match="expansion"):
        edit_pptx(source, {"a": "b" * 10_000})
    with pytest.raises(ValueError):
        create_pptx({"slides": [{"table": [["a"], ["a", "b"]]}]})


def test_actual_pdf_text_page_selection_merge_and_png_pixels():
    data = create_pdf(PDF)
    assert data.startswith(b"%PDF-")
    result = inspect_pdf(data)
    assert len(result["pages"]) == 2 and "PrismOS report" in result["pages"][0]["text"]
    assert "caf\u00e9" in result["pages"][0]["text"]
    selected = select_pdf_pages(data, [2, 1, 2])
    pages = inspect_pdf(selected)["pages"]
    assert len(pages) == 3 and "Second page" in pages[0]["text"] and "PrismOS report" in pages[1]["text"]
    merged = merge_pdf([data, select_pdf_pages(data, [1])])
    assert len(inspect_pdf(merged)["pages"]) == 3
    png = render_pdf_page(data, 1, 1)
    with Image.open(io.BytesIO(png)) as image:
        assert image.size == (595, 842)
        # Actual text must have visible pixels; a blank image is not a render.
        assert image.convert("RGB").getextrema() != ((255, 255), (255, 255), (255, 255))
    assert result["ocr_performed"] is False and result["forms_executed"] is False


@pytest.mark.parametrize("case", ["unicode-font", "line-width", "page-count", "bad-page", "scale", "bad-pdf"])
def test_pdf_unsupported_or_unbounded_operations_fail_explicitly(case):
    with pytest.raises((ValueError, UnicodeError)):
        if case == "unicode-font":
            create_pdf({"pages": [["中文字体待接入"]]})
        elif case == "line-width":
            create_pdf({"pages": [["W" * 100]]})
        elif case == "page-count":
            create_pdf({"pages": [[]] * 101})
        elif case == "bad-page":
            select_pdf_pages(create_pdf(PDF), [0])
        elif case == "scale":
            render_pdf_page(create_pdf(PDF), 1, float("nan"))
        else:
            inspect_pdf(b"not a pdf")


def test_presentation_and_pdf_native_outputs_download_through_original_routes(file_client):
    client, _, _, runtime = file_client
    ppt = office_pptx_tool.func(runtime, "create", output_path="/mnt/user-data/outputs/slides.pptx", document=SLIDES)
    pdf = office_pdf_tool.func(runtime, "create", output_path="/mnt/user-data/outputs/report.pdf", document=PDF)
    assert ppt["status"] == pdf["status"] == "ok"
    for name in ("slides.pptx", "report.pdf"):
        response = client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/" + name + "?download=true")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        assert client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/" + name, headers={"x-user": "user-b"}).status_code == 404
    rendered = office_pdf_tool.func(runtime, "render", source_path="/mnt/user-data/outputs/report.pdf", output_path="/mnt/user-data/outputs/page.png")
    assert rendered["status"] == "ok"
    png = client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/page.png")
    assert png.content.startswith(b"\x89PNG")
    assert office_pdf_tool.func(runtime, "create", output_path="/mnt/user-data/outputs/report.pdf", document=PDF)["status"] == "unavailable"


@pytest.mark.parametrize("kind", ["pptx", "pdf", "pdf-merge", "pdf-render"])
def test_current_member_withdrawal_blocks_each_format_and_multi_source_merge(owned_runtime, kind):
    factory, _, _, _, runtime = owned_runtime
    root = Path(runtime.state["thread_data"]["outputs_path"])
    (root / "source.pdf").write_bytes(create_pdf(PDF))
    (root / "source.pptx").write_bytes(create_pptx(SLIDES))

    async def withdraw():
        async with factory() as session:
            row = (await session.exec(select(TenantMembership))).scalars().one()
            row.status = "inactive"
            session.add(row)
            await session.commit()

    asyncio.run(withdraw())
    if kind == "pptx":
        result = office_pptx_tool.func(runtime, "inspect", source_path="/mnt/user-data/outputs/source.pptx")
    elif kind == "pdf-merge":
        result = office_pdf_tool.func(runtime, "merge", sources=["/mnt/user-data/outputs/source.pdf"], output_path="/mnt/user-data/outputs/merged.pdf")
    elif kind == "pdf-render":
        result = office_pdf_tool.func(runtime, "render", source_path="/mnt/user-data/outputs/source.pdf", output_path="/mnt/user-data/outputs/page.png")
    else:
        result = office_pdf_tool.func(runtime, "inspect", source_path="/mnt/user-data/outputs/source.pdf")
    assert result == {"status": "unavailable", "inspection_required": True, "retry": False}
    assert not (root / "merged.pdf").exists() and not (root / "page.png").exists()


def test_multi_source_merge_cannot_read_another_thread_or_host_file(owned_runtime):
    runtime = owned_runtime[-1]
    root = Path(runtime.state["thread_data"]["outputs_path"])
    (root / "source.pdf").write_bytes(create_pdf(PDF))
    private = root.parent / "private.pdf"
    private.write_bytes(create_pdf({"pages": [["PRIVATE OUTSIDE THREAD"]]}))
    result = office_pdf_tool.func(runtime, "merge", sources=["/mnt/user-data/outputs/source.pdf", str(private)], output_path="/mnt/user-data/outputs/merged.pdf")
    assert result["status"] == "unavailable" and not (root / "merged.pdf").exists()


def test_format_tools_are_registered_with_runtime_hidden_from_model_schema():
    from deerflow.tools.tools import BUILTIN_TOOLS

    for item in [office_pdf_tool, office_pptx_tool]:
        assert item in BUILTIN_TOOLS
        assert "runtime" not in item.tool_call_schema.model_json_schema()["properties"]


@pytest.mark.parametrize("kind", ["xlsx", "pptx", "pdf", "merge"])
def test_authority_withdrawn_during_format_processing_discards_private_result(owned_runtime, monkeypatch, kind):
    import importlib

    from test_office_docx import withdraw
    from test_office_xlsx import DOCUMENT as workbook_spec

    from deerflow.tools.builtins.office_xlsx import create_xlsx, office_xlsx_tool

    factory, _, _, _, runtime = owned_runtime
    root = Path(runtime.state["thread_data"]["outputs_path"])
    if kind == "xlsx":
        (root / "source.xlsx").write_bytes(create_xlsx(workbook_spec))
        module = importlib.import_module("deerflow.tools.builtins.office_xlsx")
        name = "inspect_xlsx"
    elif kind == "pptx":
        (root / "source.pptx").write_bytes(create_pptx(SLIDES))
        module = importlib.import_module("deerflow.tools.builtins.office_pptx")
        name = "inspect_pptx"
    else:
        (root / "source.pdf").write_bytes(create_pdf(PDF))
        module = importlib.import_module("deerflow.tools.builtins.office_pdf")
        name = "merge_pdf" if kind == "merge" else "inspect_pdf"
    original = getattr(module, name)

    def revoke(*args):
        result = original(*args)
        withdraw(factory)
        return result

    monkeypatch.setattr(module, name, revoke)
    if kind == "xlsx":
        result = office_xlsx_tool.func(runtime, "inspect", source_path="/mnt/user-data/outputs/source.xlsx")
    elif kind == "pptx":
        result = office_pptx_tool.func(runtime, "inspect", source_path="/mnt/user-data/outputs/source.pptx")
    elif kind == "merge":
        result = office_pdf_tool.func(runtime, "merge", sources=["/mnt/user-data/outputs/source.pdf"], output_path="/mnt/user-data/outputs/merged.pdf")
    else:
        result = office_pdf_tool.func(runtime, "inspect", source_path="/mnt/user-data/outputs/source.pdf")
    assert result == {"status": "unavailable", "inspection_required": True, "retry": False}
    assert not (root / "merged.pdf").exists()
