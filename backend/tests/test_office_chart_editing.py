"""Editable chart data/cache/workbooks through original current file authority."""

import copy
import importlib
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from test_file_access_boundaries import owned_runtime as owned_fixture
from test_file_access_boundaries import selected_runtime  # noqa: F401
from test_file_access_http import file_client as client_fixture
from test_office_charts import DECK
from test_office_docx import withdraw
from test_skill_member_tools import configured, scoped_client  # noqa: F401
from test_skill_reference_transactions import durable_sqlite  # noqa: F401

from deerflow.tools.builtins.office_pptx import create_pptx, edit_pptx_charts, inspect_pptx, office_pptx_tool

owned_runtime = owned_fixture
file_client = client_fixture
MODULE = importlib.import_module("deerflow.tools.builtins.office_pptx")


def specification(kind="bar"):
    deck = copy.deepcopy(DECK)
    chart = deck["slides"][0]["chart"]
    chart["kind"] = kind
    if kind == "pie":
        chart["series"] = chart["series"][:1]
    deck["slides"] += copy.deepcopy(deck["slides"])
    return deck


def replacement(kind="bar"):
    series = [{"name": "https://invalid/series", "values": [42, 50, 60]}]
    if kind != "pie":
        series += [{"name": "新系列", "values": [2, 3, 4]}, {"name": "第三系列", "values": [6, 7, 8]}]
    return {"kind": kind, "categories": ['=WEBSERVICE("https://invalid")', "第三季", "第四季"], "series": series}


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_real_chart_cache_workbook_style_position_and_other_chart_roundtrip(kind):
    source = create_pptx(specification(kind))
    deck = Presentation(io.BytesIO(source))
    shape = deck.slides[0].shapes[1]
    shape.chart.chart_style = 42
    shape.chart.series[0].format.fill.solid()
    shape.chart.series[0].format.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    shape.chart.chart_title.text_frame.paragraphs[0].font.bold = True
    stream = io.BytesIO()
    deck.save(stream)
    original = stream.getvalue()
    other = deck.slides[1].shapes[1].chart
    before_other_xml, before_other_book = other.part.blob, other.part.chart_workbook.xlsx_part.blob
    changed = edit_pptx_charts(original, [{"slide": 1, "chart": 1, "document": replacement(kind)}])
    result = Presentation(io.BytesIO(changed))
    after_shape = result.slides[0].shapes[1]
    chart = after_shape.chart
    assert list(chart.series[0].values) == [42, 50, 60]
    assert len(chart.series) == (1 if kind == "pie" else 3)
    assert [c.label for c in chart.plots[0].categories] == replacement(kind)["categories"]
    assert chart.chart_style == 42 and chart.chart_type == shape.chart.chart_type
    assert chart.series[0].format.fill.fore_color.rgb == RGBColor(0x12, 0x34, 0x56)
    assert chart.chart_title.text_frame.paragraphs[0].font.bold
    assert (after_shape.left, after_shape.top, after_shape.width, after_shape.height) == (shape.left, shape.top, shape.width, shape.height)
    untouched = result.slides[1].shapes[1].chart
    assert untouched.part.blob == before_other_xml and untouched.part.chart_workbook.xlsx_part.blob == before_other_book
    workbook = load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
    assert workbook.active["A2"].value == replacement(kind)["categories"][0] and workbook.active["A2"].data_type == "s"
    assert workbook.active["B1"].value == "https://invalid/series" and workbook.active["B1"].hyperlink is None
    assert workbook.active["B4"].value == 60
    assert list(Presentation(io.BytesIO(original)).slides[0].shapes[1].chart.series[0].values) == [12.5, 20]


def test_chart_title_can_be_changed_or_removed_and_text_stays_unchanged():
    source = create_pptx(specification())
    spec = replacement()
    spec["title"] = "更新后的图表标题"
    changed = edit_pptx_charts(source, [{"slide": 1, "chart": 1, "document": spec}])
    actual = Presentation(io.BytesIO(changed))
    assert actual.slides[0].shapes[1].chart.chart_title.text_frame.text == spec["title"]
    assert inspect_pptx(changed)["slides"][0]["texts"] == inspect_pptx(source)["slides"][0]["texts"]
    spec["title"] = ""
    removed = edit_pptx_charts(changed, [{"slide": 1, "chart": 1, "document": spec}])
    assert not Presentation(io.BytesIO(removed)).slides[0].shapes[1].chart.has_title


@pytest.mark.parametrize("mutation", ["slide", "chart", "bool", "duplicate", "kind", "external", "formula", "mismatch", "nonfinite", "oversized"])
def test_invalid_targets_and_data_fail_without_publishing_partial_edits(mutation):
    source = create_pptx(specification())
    changes = [{"slide": 1, "chart": 1, "document": replacement()}]
    item = changes[0]
    if mutation == "slide":
        item["slide"] = 3
    elif mutation == "chart":
        item["chart"] = 0
    elif mutation == "bool":
        item["slide"] = True
    elif mutation == "duplicate":
        changes *= 2
    elif mutation == "kind":
        item["document"]["kind"] = "line"
    elif mutation == "external":
        item["document"]["url"] = "https://invalid"
    elif mutation == "formula":
        item["document"]["series"][0]["values"][0] = "=SUM(B2:B3)"
    elif mutation == "mismatch":
        item["document"]["categories"] = []
    elif mutation == "nonfinite":
        item["document"]["series"][0]["values"][0] = float("nan")
    else:
        changes *= 33
    with pytest.raises(ValueError):
        edit_pptx_charts(source, changes)


@pytest.mark.parametrize("shared", ["chart", "workbook"])
def test_shared_parts_cannot_silently_change_another_chart(shared):
    deck = Presentation(io.BytesIO(create_pptx(specification())))
    first, second = [slide.shapes[1].chart for slide in deck.slides]
    if shared == "workbook":
        second.part.chart_workbook.xlsx_part = first.part.chart_workbook.xlsx_part
    else:
        slide = deck.slides[1]
        rid = slide.part.relate_to(first.part, RT.CHART)
        slide.shapes[1]._element.xpath(".//c:chart")[0].set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    stream = io.BytesIO()
    deck.save(stream)
    with pytest.raises(ValueError, match="Shared chart"):
        edit_pptx_charts(stream.getvalue(), [{"slide": 1, "chart": 1, "document": replacement()}])


def test_native_chart_edit_download_current_owner_and_exclusive_output(file_client):
    client, _, _, runtime = file_client
    source = "/mnt/user-data/outputs/source.pptx"
    assert office_pptx_tool.func(runtime, "create", output_path=source, document=specification())["status"] == "ok"
    result = office_pptx_tool.func(runtime, "edit_charts", source_path=source, output_path="/mnt/user-data/outputs/updated.pptx", chart_changes=[{"slide": 1, "chart": 1, "document": replacement()}])
    assert result["status"] == "ok"
    response = client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/updated.pptx?download=true")
    assert response.status_code == 200
    assert inspect_pptx(response.content)["slides"][0]["charts"][0]["series"][0]["values"] == [42, 50, 60]
    assert client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/updated.pptx", headers={"x-user": "user-b"}).status_code == 404
    assert office_pptx_tool.func(runtime, "edit_charts", source_path=source, output_path=source, chart_changes=[{"slide": 1, "chart": 1, "document": replacement()}])["status"] == "unavailable"


@pytest.mark.parametrize("point", ["before", "during"])
def test_member_withdrawal_cannot_publish_chart_edit(owned_runtime, monkeypatch, point):
    factory, _, _, _, runtime = owned_runtime
    root = Path(runtime.state["thread_data"]["outputs_path"])
    original = create_pptx(specification())
    (root / "source.pptx").write_bytes(original)
    if point == "before":
        withdraw(factory)
    else:
        edit = MODULE.edit_pptx_charts

        def revoke(*args):
            result = edit(*args)
            withdraw(factory)
            return result

        monkeypatch.setattr(MODULE, "edit_pptx_charts", revoke)
    result = office_pptx_tool.func(runtime, "edit_charts", source_path="/mnt/user-data/outputs/source.pptx", output_path="/mnt/user-data/outputs/revoked.pptx", chart_changes=[{"slide": 1, "chart": 1, "document": replacement()}])
    assert result == {"status": "unavailable", "inspection_required": True, "retry": False}
    assert not (root / "revoked.pptx").exists() and (root / "source.pptx").read_bytes() == original


def test_chart_edit_schema_keeps_runtime_hidden():
    properties = office_pptx_tool.tool_call_schema.model_json_schema()["properties"]
    assert "chart_changes" in properties and "runtime" not in properties
