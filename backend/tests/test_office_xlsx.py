"""Actual independent workbook roundtrips through original current file authority."""

import asyncio
import importlib
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_file_access_boundaries import owned_runtime as owned_fixture
from test_file_access_boundaries import selected_runtime  # noqa: F401
from test_file_access_http import file_client as client_fixture
from test_skill_member_tools import configured, scoped_client  # noqa: F401
from test_skill_reference_transactions import durable_sqlite  # noqa: F401

from deerflow.database.models import TenantMembership
from deerflow.tools.builtins.office_xlsx import create_xlsx, edit_xlsx, inspect_xlsx, office_xlsx_tool

owned_runtime = owned_fixture
file_client = client_fixture
module = importlib.import_module("deerflow.tools.builtins.office_xlsx")
OUTPUT = "/mnt/user-data/outputs/report.xlsx"
DOCUMENT = {
    "sheets": [{"name": "预算", "header": True, "filter": True, "rows": [["部门", "预算"], ["研发", 12.5], ["市场", 20], ["总计", {"formula": "=SUM(B2:B3)"}]]}, {"name": "说明", "rows": [['=WEBSERVICE("https://invalid")', "中文"]]}]
}


def test_actual_workbook_and_independent_reader_cell_types():
    data = create_xlsx(DOCUMENT)
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert workbook.sheetnames == ["预算", "说明"]
    assert workbook["预算"]["B2"].value == 12.5
    assert workbook["预算"]["B4"].data_type == "f" and workbook["预算"]["B4"].value == "=SUM(B2:B3)"
    assert workbook["说明"]["A1"].data_type == "s"
    assert workbook["预算"]["A1"].font.bold and workbook["预算"].freeze_panes == "A2"
    assert workbook["预算"].auto_filter.ref == "A1:B4"
    result = inspect_xlsx(data)
    assert result["sheets"][0]["rows"][3][1] == {"formula": "=SUM(B2:B3)", "calculated": False}
    assert result["formulas_calculated"] is False
    assert load_workbook(io.BytesIO(data), data_only=True)["预算"]["B4"].value is None


def test_edit_preserves_other_sheets_styles_merges_and_original_bytes():
    data = create_xlsx(DOCUMENT)
    workbook = load_workbook(io.BytesIO(data))
    workbook["说明"].merge_cells("B2:C2")
    output = io.BytesIO()
    workbook.save(output)
    original = output.getvalue()
    changed = edit_xlsx(original, [{"sheet": "预算", "cell": "B2", "value": 30}, {"sheet": "说明", "cell": "D3", "value": "新说明"}])
    result = load_workbook(io.BytesIO(changed), data_only=False)
    assert result["预算"]["B2"].value == 30 and result["预算"]["A1"].font.bold
    assert "B2:C2" in result["说明"].merged_cells
    assert result["说明"]["D3"].value == "新说明"
    assert load_workbook(io.BytesIO(original))["预算"]["B2"].value == 12.5


@pytest.mark.parametrize("formula", ['=WEBSERVICE("https://invalid")', '=HYPERLINK("file:///private")', "='[foreign.xlsx]Sheet'!A1", "=cmd|'payload'!A1", "=_xlfn.UNKNOWN(A1)"])
def test_active_or_external_formulas_are_not_generated(formula):
    with pytest.raises(ValueError):
        create_xlsx({"sheets": [{"name": "Sheet", "rows": [[{"formula": formula}]]}]})


def test_native_workbook_download_edit_and_exclusive_result(file_client):
    client, _, _, runtime = file_client
    created = office_xlsx_tool.func(runtime, "create", output_path=OUTPUT, document=DOCUMENT)
    assert created["status"] == "ok"
    response = client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/report.xlsx?download=true")
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    assert load_workbook(io.BytesIO(response.content))["预算"]["B2"].value == 12.5
    assert office_xlsx_tool.func(runtime, "create", output_path=OUTPUT, document=DOCUMENT)["status"] == "unavailable"
    assert office_xlsx_tool.func(runtime, "inspect", source_path=OUTPUT)["status"] == "ok"
    edited = office_xlsx_tool.func(runtime, "edit", source_path=OUTPUT, output_path=OUTPUT.replace("report", "edited"), changes=[{"sheet": "预算", "cell": "B2", "value": 42}])
    assert edited["status"] == "ok"
    assert client.get("/api/threads/member-thread/artifacts/mnt/user-data/outputs/report.xlsx", headers={"x-user": "user-b"}).status_code == 404


@pytest.mark.parametrize("action", ["create", "inspect", "edit"])
def test_current_member_revocation_blocks_workbook_operation(owned_runtime, action):
    factory, _, _, _, runtime = owned_runtime
    output = Path(runtime.state["thread_data"]["outputs_path"])
    (output / "source.xlsx").write_bytes(create_xlsx(DOCUMENT))

    async def withdraw():
        async with factory() as session:
            row = (await session.exec(select(TenantMembership))).scalars().one()
            row.status = "inactive"
            session.add(row)
            await session.commit()

    asyncio.run(withdraw())
    result = office_xlsx_tool.func(runtime, action, source_path="/mnt/user-data/outputs/source.xlsx", output_path=OUTPUT, document=DOCUMENT, changes=[])
    assert result == {"status": "unavailable", "inspection_required": True, "retry": False}
    assert not (output / "report.xlsx").exists()


@pytest.mark.parametrize("path", ["/etc/report.xlsx", "/mnt/user-data/outputs/../uploads/report.xlsx", "/mnt/user-data/uploads/report.xlsx", "/mnt/user-data/outputs/fake.xlsm"])
def test_output_path_and_format_cannot_be_forged(owned_runtime, path):
    runtime = owned_runtime[-1]
    assert office_xlsx_tool.func(runtime, "create", output_path=path, document=DOCUMENT)["status"] == "unavailable"
    assert not list(Path(runtime.state["thread_data"]["uploads_path"]).iterdir())


@pytest.mark.parametrize("kind", ["external", "macro", "duplicate", "entity", "depth"])
def test_workbook_package_rejects_active_and_oversized_structure(kind):
    with zipfile.ZipFile(io.BytesIO(create_xlsx(DOCUMENT))) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    if kind == "external":
        parts["xl/_rels/workbook.xml.rels"] = b'<Relationships><Relationship TargetMode="External" Target="file:///private"/></Relationships>'
    elif kind == "macro":
        parts["xl/vbaProject.bin"] = b"unavailable"
    elif kind == "entity":
        parts["docProps/custom.xml"] = '<!DOCTYPE x [<!ENTITY a SYSTEM "file:///private">]><x>&a;</x>'.encode("utf-16")
    elif kind == "depth":
        parts["docProps/custom.xml"] = ("<x>" * 102 + "a" + "</x>" * 102).encode()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
        if kind == "duplicate":
            with pytest.warns(UserWarning):
                archive.writestr("xl/workbook.xml", parts["xl/workbook.xml"])
    with pytest.raises(ValueError):
        inspect_xlsx(output.getvalue())


def test_cell_addresses_and_workbook_dimensions_are_bounded():
    source = create_xlsx(DOCUMENT)
    for address in ["A0", "A10001", "ZZ100", "../A1", "A1:B2"]:
        with pytest.raises(ValueError):
            edit_xlsx(source, [{"sheet": "预算", "cell": address, "value": "no"}])
    with pytest.raises(ValueError):
        create_xlsx({"sheets": [{"name": "Sheet", "rows": [["a"] * 257]}]})


def test_workbook_tool_schema_and_registration():
    from deerflow.tools.tools import BUILTIN_TOOLS

    assert office_xlsx_tool in BUILTIN_TOOLS
    assert "runtime" not in office_xlsx_tool.tool_call_schema.model_json_schema()["properties"]
