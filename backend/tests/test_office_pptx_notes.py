"""Speaker-note editing preserves editable slides/charts under current authority."""

import io
from pathlib import Path

import pytest
from pptx import Presentation
from test_office_chart_editing import configured, durable_sqlite, scoped_client, selected_runtime  # noqa: F401
from test_office_chart_editing import file_client as client_fixture
from test_office_chart_editing import owned_runtime as owned_fixture
from test_office_charts import DECK
from test_office_docx import withdraw

from deerflow.tools.builtins.office_pptx import create_pptx, edit_pptx_notes, inspect_pptx, office_pptx_tool

file_client = client_fixture
owned_runtime = owned_fixture


def test_add_replace_clear_notes_preserves_slide_chart_and_embedded_data():
    source = create_pptx(DECK)
    original = Presentation(io.BytesIO(source))
    chart = original.slides[0].shapes[1].chart
    notes = edit_pptx_notes(source, [{"slide": 1, "text": "原始数据来源\n中文演讲者备注"}])
    actual = Presentation(io.BytesIO(notes))
    after = actual.slides[0].shapes[1].chart
    assert after.part.blob == chart.part.blob
    assert after.part.chart_workbook.xlsx_part.blob == chart.part.chart_workbook.xlsx_part.blob
    assert actual.slides[0]._element.xml == original.slides[0]._element.xml
    assert inspect_pptx(notes)["slides"][0]["notes"] == "原始数据来源\n中文演讲者备注"
    replaced = edit_pptx_notes(notes, [{"slide": 1, "text": "更正收入175"}])
    assert inspect_pptx(replaced)["slides"][0]["notes"] == "更正收入175"
    assert inspect_pptx(edit_pptx_notes(replaced, [{"slide": 1, "text": ""}]))["slides"][0]["notes"] == ""
    assert inspect_pptx(source)["slides"][0]["notes"] == DECK["slides"][0].get("notes", "")


@pytest.mark.parametrize(
    "notes",
    [
        [],
        [{"slide": True, "text": "bad"}],
        [{"slide": 0, "text": "bad"}],
        [{"slide": 999, "text": "bad"}],
        [{"slide": 1, "text": "a"}, {"slide": 1, "text": "b"}],
        [{"slide": 1, "text": "bad\x01"}],
        [{"slide": 1, "text": "bad", "other": 1}],
    ],
)
def test_notes_reject_invalid_targets_or_controls(notes):
    with pytest.raises(ValueError):
        edit_pptx_notes(create_pptx(DECK), notes)


def test_notes_use_current_authority_and_original_download(file_client):
    client, factory, _, runtime = file_client
    root = Path(runtime.state["thread_data"]["outputs_path"])
    source = root / "notes-source.pptx"
    original = create_pptx(DECK)
    source.write_bytes(original)
    result = office_pptx_tool.func(runtime, "edit_notes", source_path="/mnt/user-data/outputs/notes-source.pptx", output_path="/mnt/user-data/outputs/notes-edited.pptx", notes=[{"slide": 1, "text": "当前有权备注"}])
    assert result["status"] == "ok"
    assert source.read_bytes() == original
    response = client.get("/api/threads/" + runtime.context["thread_id"] + "/artifacts/mnt/user-data/outputs/notes-edited.pptx?download=true")
    assert response.status_code == 200
    assert inspect_pptx(response.content)["slides"][0]["notes"] == "当前有权备注"
    withdraw(factory)
    denied = office_pptx_tool.func(runtime, "edit_notes", source_path="/mnt/user-data/outputs/notes-source.pptx", output_path="/mnt/user-data/outputs/notes-denied.pptx", notes=[{"slide": 1, "text": "revoked"}])
    assert denied["status"] == "unavailable" and not (root / "notes-denied.pptx").exists()


def test_model_schema_exposes_slide_notes_and_note_updates():
    schema = office_pptx_tool.tool_call_schema.model_json_schema()
    assert "notes" in schema["$defs"]["PptxSlide"]["properties"]
    assert set(schema["$defs"]["PptxNoteChange"]["required"]) == {"slide", "text"}
    assert schema["$defs"]["PptxChart"]["properties"]["kind"]["enum"] == ["bar", "line", "pie"]


@pytest.mark.parametrize("value", [True, "175"])
def test_typed_chart_numeric_schema_does_not_coerce_invalid_data(value):
    from pydantic import ValidationError

    document = {"slides": [{"chart": {"kind": "bar", "categories": ["week"], "series": [{"name": "revenue", "values": [value]}]}}]}
    with pytest.raises(ValidationError):
        office_pptx_tool.tool_call_schema.model_validate({"action": "create", "document": document})
    with pytest.raises(ValidationError):
        office_pptx_tool.tool_call_schema.model_validate({"action": "edit_notes", "notes": [{"slide": value, "text": "invalid"}]})
