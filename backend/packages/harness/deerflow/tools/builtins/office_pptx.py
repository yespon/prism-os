"""Independent bounded presentations; no copied upstream Office scripts."""

import io
import json
import math
import zipfile

from langchain.tools import ToolRuntime, tool
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from .office_common import check_office_package, owned_office_operation
from .office_xlsx import _budget as _workbook_budget
from .office_xlsx import _load as _load_workbook
from .office_xlsx import create_xlsx

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def _text(value):
    if not isinstance(value, str) or len(value) > 100_000 or any(ord(c) < 32 and c not in "\t\n\r" for c in value):
        raise ValueError("Presentation text unavailable")
    return value


def _shapes(collection, depth=0):
    if depth > 16:
        raise ValueError("Presentation group depth exceeded")
    for shape in collection:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes, depth + 1)


def _frames(slide):
    count = 0
    for shape in _shapes(slide.shapes):
        count += 1
        if count > 2000:
            raise ValueError("Presentation shape limit exceeded")
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        yield slide.notes_slide.notes_text_frame


def _budget(presentation):
    if len(presentation.slides) > 100:
        raise ValueError("Presentation slide limit exceeded")
    total = charts = 0
    for slide in presentation.slides:
        for shape in _shapes(slide.shapes):
            if shape.has_chart:
                charts += 1
                chart = shape.chart
                if charts > 32 or len(chart.series) > 8 or not chart.plots or any(len(plot.categories) > 200 for plot in chart.plots):
                    raise ValueError("Presentation chart limit exceeded")
        for frame in _frames(slide):
            total += len(frame.text)
            if total > 2_000_000:
                raise ValueError("Presentation text limit exceeded")


def _embedded_workbooks(data):
    total = count = 0
    with zipfile.ZipFile(io.BytesIO(data)) as outer:
        for name in outer.namelist():
            if not name.startswith("ppt/embeddings/") or name.endswith("/"):
                continue
            count += 1
            if not name.endswith(".xlsx") or count > 32:
                raise ValueError("Presentation embedded content unavailable")
            body = outer.read(name)
            with zipfile.ZipFile(io.BytesIO(body)) as workbook:
                total += sum(item.file_size for item in workbook.infolist())
            if total > 8 * 1024 * 1024:
                raise ValueError("Presentation embedded workbook limit exceeded")
            _workbook_budget(_load_workbook(body), validate_formulas=True)


def _chart_data(spec):
    if not isinstance(spec, dict) or set(spec) - {"kind", "title", "categories", "series"} or spec.get("kind") not in {"bar", "line", "pie"}:
        raise ValueError("Presentation chart schema unavailable")
    categories, series = spec.get("categories"), spec.get("series")
    if not isinstance(categories, list) or not 1 <= len(categories) <= 200 or not isinstance(series, list) or not 1 <= len(series) <= 8:
        raise ValueError("Presentation chart dimensions unavailable")
    if any(not isinstance(value, str) or not value or len(value) > 1000 for value in categories):
        raise ValueError("Presentation chart category unavailable")
    if spec["kind"] == "pie" and len(series) != 1:
        raise ValueError("Pie chart requires one series")
    data = CategoryChartData()
    data.categories = [_text(value) for value in categories]
    for item in series:
        if not isinstance(item, dict) or set(item) != {"name", "values"} or not isinstance(item["name"], str) or not item["name"] or len(item["name"]) > 1000:
            raise ValueError("Presentation chart series unavailable")
        values = item["values"]
        if not isinstance(values, list) or len(values) != len(categories) or any(type(value) not in {int, float} or not math.isfinite(value) or (spec["kind"] == "pie" and value < 0) for value in values):
            raise ValueError("Presentation chart requires finite matching numeric values")
        data.add_series(_text(item["name"]), values)
    title = spec.get("title", "")
    if not isinstance(title, str) or len(title) > 1000:
        raise ValueError("Presentation chart title unavailable")
    return data


def _chart_workbook(spec):
    categories, series = spec["categories"], spec["series"]
    rows = [["", *[item["name"] for item in series]]]
    rows += [[category, *[item["values"][index] for item in series]] for index, category in enumerate(categories)]
    return create_xlsx({"sheets": [{"name": "Sheet1", "rows": rows}]})


def _chart(slide, spec):
    data = _chart_data(spec)
    series = spec["series"]
    title = spec.get("title", "")
    chart = slide.shapes.add_chart({"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}[spec["kind"]], Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.1), data).chart
    # The chart engine normally interprets '=' and URL-like labels in its XLSX
    # writer. Replace its data part with our literal-cell workbook before output.
    chart.part.chart_workbook.update_from_xlsx_blob(_chart_workbook(spec))
    chart.chart_style = 10
    chart.has_legend = len(series) > 1 or spec["kind"] == "pie"
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = _text(title)


def _load(data):
    check_office_package(data, "ppt/presentation.xml", CONTENT_TYPE)
    _embedded_workbooks(data)
    result = Presentation(io.BytesIO(data))
    _budget(result)
    return result


def _save(presentation):
    _budget(presentation)
    stream = io.BytesIO()
    presentation.save(stream)
    data = stream.getvalue()
    check_office_package(data, "ppt/presentation.xml", CONTENT_TYPE)
    _embedded_workbooks(data)
    return data


def create_pptx(document):
    if not isinstance(document, dict) or set(document) != {"slides"} or len(json.dumps(document, ensure_ascii=False)) > 2_000_000:
        raise ValueError("Presentation schema unavailable")
    slides = document["slides"]
    if not isinstance(slides, list) or not 1 <= len(slides) <= 100:
        raise ValueError("Presentation slides unavailable")
    presentation = Presentation()
    # The engine default contains an unreferenced printer configuration binary.
    # Remove that relationship from new documents; arbitrary input binaries remain denied.
    for rel in list(presentation.part.rels.values()):
        if rel.reltype == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/printerSettings":
            if any(rel.rId in node.attrib.values() for node in presentation._element.iter()):
                raise ValueError("Referenced printer configuration unavailable")
            presentation.part.drop_rel(rel.rId)
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    chart_count = 0
    for item in slides:
        if not isinstance(item, dict) or set(item) - {"title", "paragraphs", "table", "notes", "chart"} or sum(key in item for key in ["table", "paragraphs", "chart"]) > 1:
            raise ValueError("Presentation slide schema unavailable")
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9)).text_frame
        title.text = _text(item.get("title", ""))
        title.paragraphs[0].font.size = Pt(32)
        title.paragraphs[0].font.bold = True
        if "chart" in item:
            chart_count += 1
            if chart_count > 32:
                raise ValueError("Presentation chart limit exceeded")
            _chart(slide, item["chart"])
        elif "table" in item:
            rows = item["table"]
            if not isinstance(rows, list) or not 1 <= len(rows) <= 50 or not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 20:
                raise ValueError("Presentation table unavailable")
            columns = len(rows[0])
            if any(not isinstance(row, list) or len(row) != columns for row in rows):
                raise ValueError("Presentation rectangular table required")
            table = slide.shapes.add_table(len(rows), columns, Inches(0.7), Inches(1.6), Inches(11.9), Inches(4.8)).table
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    table.cell(i, j).text = _text(value)
        else:
            values = item.get("paragraphs", [])
            if not isinstance(values, list) or len(values) > 100:
                raise ValueError("Presentation paragraphs unavailable")
            frame = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.1)).text_frame
            frame.word_wrap = True
            for i, value in enumerate(values):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text = _text(value)
                paragraph.font.size = Pt(22)
        if "notes" in item:
            slide.notes_slide.notes_text_frame.text = _text(item["notes"])
    return _save(presentation)


def inspect_pptx(data):
    presentation = _load(data)
    slides = []
    for slide in presentation.slides:
        texts, tables, charts = [], [], []
        for shape in _shapes(slide.shapes):
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_chart:
                chart = shape.chart
                charts.append({"kind": str(chart.chart_type), "series": [{"name": item.name, "values": list(item.values)} for item in chart.series], "categories": [item.label for item in chart.plots[0].categories]})
            if shape.has_table:
                tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
        slides.append({"texts": texts, "tables": tables, "charts": charts, "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ""})
    return {"slides": slides, "size_emu": [presentation.slide_width, presentation.slide_height], "rendered": False}


def edit_pptx(data, replacements):
    if not isinstance(replacements, dict) or not replacements or len(replacements) > 100:
        raise ValueError("Presentation replacements unavailable")
    for old, new in replacements.items():
        if not _text(old):
            raise ValueError("Empty replacement unavailable")
        _text(new)
    presentation = _load(data)
    matches = total = 0
    for slide in presentation.slides:
        for frame in _frames(slide):
            for paragraph in frame.paragraphs:
                current = paragraph.text
                for old, new in replacements.items():
                    count = current.count(old)
                    if len(current) + count * (len(new) - len(old)) > 100_000:
                        raise ValueError("Presentation replacement expansion limit exceeded")
                    matches += count
                    current = current.replace(old, new)
                total += len(current)
                if total > 2_000_000:
                    raise ValueError("Presentation text limit exceeded")
                if current != paragraph.text:
                    runs = list(paragraph.runs)
                    if runs:
                        runs[0].text = current
                        for run in runs[1:]:
                            run.text = ""
                    else:
                        paragraph.text = current
    if not matches:
        raise ValueError("No presentation replacement matched")
    return _save(presentation)


def edit_pptx_charts(data, changes):
    if not isinstance(changes, list) or not 1 <= len(changes) <= 32 or len(json.dumps(changes, ensure_ascii=False)) > 2_000_000:
        raise ValueError("Presentation chart changes unavailable")
    presentation = _load(data)
    all_charts = [shape.chart for slide in presentation.slides for shape in _shapes(slide.shapes) if shape.has_chart]
    seen = set()
    for change in changes:
        if not isinstance(change, dict) or set(change) != {"slide", "chart", "document"} or type(change["slide"]) is not int or type(change["chart"]) is not int:
            raise ValueError("Presentation chart target unavailable")
        slide_number, chart_number = change["slide"], change["chart"]
        if not 1 <= slide_number <= len(presentation.slides):
            raise ValueError("Presentation chart slide unavailable")
        charts = [shape.chart for shape in _shapes(presentation.slides[slide_number - 1].shapes) if shape.has_chart]
        if not 1 <= chart_number <= len(charts) or (slide_number, chart_number) in seen:
            raise ValueError("Presentation chart target unavailable")
        seen.add((slide_number, chart_number))
        chart = charts[chart_number - 1]
        workbook_part = chart.part.chart_workbook.xlsx_part
        if sum(other.part is chart.part for other in all_charts) != 1 or (workbook_part is not None and sum(other.part.chart_workbook.xlsx_part is workbook_part for other in all_charts) != 1):
            raise ValueError("Shared chart data requires an independent copy")
        kind = {XL_CHART_TYPE.COLUMN_CLUSTERED: "bar", XL_CHART_TYPE.LINE: "line", XL_CHART_TYPE.PIE: "pie"}.get(chart.chart_type)
        document = change["document"]
        if kind is None or len(chart.plots) != 1 or not isinstance(document, dict) or document.get("kind") != kind:
            raise ValueError("Existing chart kind must be preserved")
        replacement = _chart_data(document)
        chart.replace_data(replacement)
        chart.part.chart_workbook.update_from_xlsx_blob(_chart_workbook(document))
        if "title" in document:
            title = _text(document["title"])
            chart.has_title = bool(title)
            if title:
                chart.chart_title.text_frame.text = title
    return _save(presentation)


@tool("office_pptx", parse_docstring=True)
def office_pptx_tool(runtime: ToolRuntime, action: str, source_path: str = "", output_path: str = "", document: dict | None = None, replacements: dict[str, str] | None = None, chart_changes: list[dict] | None = None) -> dict:
    """Create, inspect or edit an owned PPTX without rendering slides or executing content.

    Args:
        action: create, inspect, edit (literal text) or edit_charts (chart data). Desktop rendering, animations and advanced layouts remain pending.
        source_path: Current authorized PPTX for inspect or edit.
        output_path: New /mnt/user-data/outputs/*.pptx, never overwrites; present with present_files.
        document: Create schema: slides list of title, paragraphs (strings) OR table (rectangular strings) OR chart (kind bar/line/pie, optional title, categories strings, series
            list of name and finite numeric values), optional notes. Charts embed editable XLSX data; max 200 categories and 8 series, pie one nonnegative series. Uses a simple
            wide layout.
        replacements: Literal text replacements in slide text, tables and notes. Cross-run matches adopt the first run style; advanced roundtrip fidelity is not claimed.
        chart_changes: edit_charts list of slide/chart (1-based chart order including groups), document (same chart create schema). Keeps chart kind/position/style; title omitted
            preserves it, empty title removes it. Shared chart/workbook parts require an independent copy.
    """
    edit = (lambda data: edit_pptx_charts(data, chart_changes)) if action == "edit_charts" else (lambda data: edit_pptx(data, replacements))
    mapped = "edit" if action == "edit_charts" else action
    return owned_office_operation(runtime, mapped, ".pptx", source_path, output_path, lambda: create_pptx(document), inspect_pptx, edit)
