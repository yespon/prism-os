"""Independent owned PNG/JPEG insertion, with finite slide geometry and current ACL."""

import io
import json
import math

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from deerflow.runtime import current_file_access

from .office_docx_images import _image


def add_pptx_images(data, images):
    from .office_pptx import _load, _save, _shapes, _text

    if not isinstance(images, list) or not 1 <= len(images) <= 16:
        raise ValueError('Presentation images unavailable')
    presentation = _load(data)
    count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in presentation.slides for shape in _shapes(slide.shapes))
    if count + len(images) > 64:
        raise ValueError('Presentation picture limit exceeded')
    total = 0
    for item in images:
        if not isinstance(item, dict) or set(item) - {'slide', 'data', 'left_inches', 'top_inches', 'width_inches', 'alt'}:
            raise ValueError('Presentation image schema unavailable')
        index = item.get('slide')
        if type(index) is not int or not 1 <= index <= len(presentation.slides):
            raise ValueError('Presentation image slide unavailable')
        image, width, height = _image(item.get('data'))
        total += len(image)
        if total > 8 * 1024 * 1024:
            raise ValueError('Presentation image total unavailable')
        left, top, size = item.get('left_inches', 0.7), item.get('top_inches', 1.6), item.get('width_inches', 6)
        if any(type(value) not in {int, float} or not math.isfinite(value) for value in [left, top, size]) or min(left, top) < 0 or size < 0.1:
            raise ValueError('Presentation image geometry unavailable')
        cy = size * height / width
        slide_width = presentation.slide_width / 914400
        slide_height = presentation.slide_height / 914400
        if left + size > slide_width or top + cy > slide_height:
            raise ValueError('Presentation image outside slide')
        alt = item.get('alt', '')
        if not isinstance(alt, str) or len(alt) > 2000:
            raise ValueError('Presentation image description unavailable')
        _text(alt)
        slide = presentation.slides[index - 1]
        picture = slide.shapes.add_picture(io.BytesIO(image), Inches(left), Inches(top), width=Inches(size), height=Inches(cy))
        picture._element.nvPicPr.cNvPr.set('descr', alt)
    return _save(presentation)


def add_owned_pptx_images(runtime, data, images):
    if not isinstance(images, list) or not 1 <= len(images) <= 16 or len(json.dumps(images, ensure_ascii=False)) > 100_000:
        raise ValueError('Owned presentation images unavailable')
    resolved, total = [], 0
    for item in images:
        if not isinstance(item, dict) or set(item) - {'slide', 'source_path', 'left_inches', 'top_inches', 'width_inches', 'alt'} or not isinstance(item.get('source_path'), str) or not item['source_path'].lower().endswith(
            ('.png', '.jpg', '.jpeg')
        ):
            raise ValueError('Owned PNG/JPEG image required')
        image = current_file_access(runtime, 'read_bytes', item['source_path'])
        if not isinstance(image, bytes) or len(image) > 4 * 1024 * 1024:
            raise ValueError('Owned image size unavailable')
        total += len(image)
        if total > 8 * 1024 * 1024:
            raise ValueError('Owned image total unavailable')
        resolved.append({'data': image, **{key: value for key, value in item.items() if key != 'source_path'}})
    result = add_pptx_images(data, resolved)
    for item in images:
        current_file_access(runtime, 'read_bytes', item['source_path'])
    return result
