"""Independent spreadsheet operations using the explicit locked openpyxl engine."""

import io
import json
import math
import re

from langchain.tools import ToolRuntime, tool
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import coordinate_to_tuple

from .office_common import check_office_package, owned_office_operation

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
SAFE_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "SUMIF", "COUNTIF", "ROUND", "ABS", "AND", "OR", "NOT"}


def _formula(value):
    if not isinstance(value, str) or not value.startswith("=") or len(value) > 2000 or any(c in value for c in "[]|\r\n"):
        raise ValueError("Spreadsheet formula unavailable")
    functions = re.findall(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(", value)
    if any(name.upper() not in SAFE_FUNCTIONS for name in functions) or re.search(r"https?:|file:|\\", value, re.I):
        raise ValueError("Spreadsheet formula unavailable")
    return value


def _set(cell, value):
    if isinstance(value, dict) and set(value) == {"formula"}:
        cell.value = _formula(value["formula"])
    elif value is None or isinstance(value, (str, bool, int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("Finite spreadsheet number required")
        if isinstance(value, str) and (len(value) > 32767 or any(ord(c) < 32 and c not in "\r\n\t" for c in value)):
            raise ValueError("Spreadsheet text unavailable")
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"  # Plain input never becomes an Excel formula.
    else:
        raise ValueError("Spreadsheet cell unavailable")


def _budget(workbook, *, validate_formulas=False):
    cells = text = 0
    if len(workbook.worksheets) > 32:
        raise ValueError("Spreadsheet sheet limit exceeded")
    for sheet in workbook:
        if sheet.max_row > 10_000 or sheet.max_column > 256:
            raise ValueError("Spreadsheet dimensions limit exceeded")
        for row in sheet.iter_rows():
            for cell in row:
                cells += 1
                text += len(str(cell.value)) if cell.value is not None else 0
                if cells > 100_000 or text > 2_000_000:
                    raise ValueError("Spreadsheet content limit exceeded")
                if validate_formulas and cell.data_type == "f":
                    _formula(cell.value)


def _load(data):
    check_office_package(data, "xl/workbook.xml", CONTENT_TYPE)
    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
    _budget(workbook)
    return workbook


def _save(workbook):
    _budget(workbook, validate_formulas=True)
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    check_office_package(data, "xl/workbook.xml", CONTENT_TYPE)
    return data


def create_xlsx(document):
    if not isinstance(document, dict) or set(document) != {"sheets"} or len(json.dumps(document, ensure_ascii=False)) > 2_000_000:
        raise ValueError("Spreadsheet schema unavailable")
    sheets = document["sheets"]
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 32:
        raise ValueError("Spreadsheet sheets unavailable")
    workbook = Workbook()
    workbook.remove(workbook.active)
    names = set()
    cells = 0
    for item in sheets:
        if not isinstance(item, dict) or set(item) - {"name", "rows", "header", "filter"}:
            raise ValueError("Spreadsheet sheet schema unavailable")
        name, rows = item.get("name"), item.get("rows", [])
        if not isinstance(name, str) or not name or len(name) > 31 or name.casefold() in names or not isinstance(rows, list) or len(rows) > 10_000:
            raise ValueError("Spreadsheet sheet unavailable")
        names.add(name.casefold())
        sheet = workbook.create_sheet(name)
        for i, values in enumerate(rows, 1):
            if not isinstance(values, list) or len(values) > 256:
                raise ValueError("Spreadsheet row unavailable")
            cells += len(values)
            if cells > 100_000:
                raise ValueError("Spreadsheet cell limit exceeded")
            for j, value in enumerate(values, 1):
                _set(sheet.cell(i, j), value)
        if item.get("header") and rows:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
        if item.get("filter") and rows:
            sheet.auto_filter.ref = sheet.dimensions
    return _save(workbook)


def inspect_xlsx(data):
    workbook = _load(data)
    result = []
    for sheet in workbook:
        rows = []
        for row in sheet.iter_rows():
            rows.append([{"formula": cell.value, "calculated": False} if cell.data_type == "f" else cell.value.isoformat() if hasattr(cell.value, "isoformat") else cell.value for cell in row])
        result.append({"name": sheet.title, "rows": rows, "merged_cells": [str(value) for value in sheet.merged_cells.ranges]})
    return {"sheets": result, "formulas_calculated": False}


def edit_xlsx(data, changes):
    if not isinstance(changes, list) or len(changes) > 10_000 or len(json.dumps(changes, ensure_ascii=False)) > 2_000_000:
        raise ValueError("Spreadsheet changes unavailable")
    workbook = _load(data)
    for change in changes:
        if not isinstance(change, dict) or set(change) != {"sheet", "cell", "value"} or change["sheet"] not in workbook.sheetnames:
            raise ValueError("Spreadsheet change unavailable")
        address = change["cell"]
        if not isinstance(address, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,4}", address):
            raise ValueError("Spreadsheet address unavailable")
        row, column = coordinate_to_tuple(address)
        if row > 10_000 or column > 256:
            raise ValueError("Spreadsheet dimensions limit exceeded")
        _set(workbook[change["sheet"]].cell(row, column), change["value"])
    return _save(workbook)


@tool("office_xlsx", parse_docstring=True)
def office_xlsx_tool(runtime: ToolRuntime, action: str, source_path: str = "", output_path: str = "", document: dict | None = None, changes: list[dict] | None = None) -> dict:
    """Create, inspect or edit owned XLSX without evaluating formulas or executing macros.

    Args:
        action: create, inspect or edit. Advanced pivot/chart/calculate/render operations remain pending.
        source_path: Current authorized XLSX input for inspect or edit.
        output_path: New /mnt/user-data/outputs/*.xlsx, never overwrites. Present with present_files.
        document: Create schema: sheets list of name, rows (cell values), optional header/filter booleans. Strings are literal; explicit {formula: '=SUM(A1:A3)'} supports a limited safe formula set, not calculated values.
        changes: Edit list of sheet, cell (A1 address) and value. Unchanged worksheets/styles/merges are retained by the workbook engine; full Office roundtrip fidelity is not claimed.
    """
    return owned_office_operation(runtime, action, ".xlsx", source_path, output_path, lambda: create_xlsx(document), inspect_xlsx, lambda data: edit_xlsx(data, changes))
