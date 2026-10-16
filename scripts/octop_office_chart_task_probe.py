"""Actual original Office expert creates and revises editable chart deliverables."""

import argparse
import hashlib
import io
import json
import secrets
from pathlib import Path

import httpx
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
STATE = ROOT / "backend/tmp/n04-wsl"
BASE = "http://127.0.0.1:18026"


def tool_trace(transcript):
    calls = {}
    for line in transcript.splitlines():
        if not line.startswith("data:"):
            continue
        try:
            stack = [json.loads(line[5:])]
        except ValueError:
            continue
        while stack:
            value = stack.pop()
            if isinstance(value, list):
                stack.extend(value)
            elif isinstance(value, dict):
                for call in value.get("tool_calls", []):
                    if call.get("id") and call.get("name"):
                        calls[call["id"]] = {
                            "id": call["id"],
                            "name": call["name"],
                            "action": call.get("args", {}).get("action"),
                        }
                stack.extend(v for v in value.values() if isinstance(v, (dict, list)))
    return list(calls.values())


def execute(output, resume=None):
    assert not output.exists(), "Preserve previous receipts"
    fixture = json.loads((STATE / "live-fixture.json").read_text())
    previous = json.loads(resume.read_text()) if resume else None
    suffix = previous["suffix"] if previous else secrets.token_hex(5)
    private = STATE / ("office-charts-" + secrets.token_hex(5))
    private.mkdir(mode=0o700)
    result = dict(previous) if previous else {"suffix": suffix, "phases": []}
    if previous:
        result.pop("failure", None)
        result.pop("failure_type", None)
        result["previous_receipt"] = str(resume)
        if not result.get("actual_tool_calls"):
            source_private = Path(previous["private_diagnostics"])
            result["actual_tool_calls"] = tool_trace(
                (source_private / "initial-stream.txt").read_text()
            )
            first_state = json.loads(
                (source_private / "initial-state.json").read_text()
            )
            result["first_task_checkpoint"] = first_state["checkpoint"]
            result["first_task_message_ids"] = [
                m.get("id") for m in first_state["values"]["messages"]
            ]
    result.update(
        complete=False,
        private_diagnostics=str(private),
        actual_model=fixture["model_name"],
        simulated_model=False,
        new_model_task_submissions=0,
    )
    clients = []

    def save():
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")

    def call(client, path, body=None, status=200, method=None):
        response = client.request(
            method or ("POST" if body is not None else "GET"), path, json=body
        )
        assert response.status_code == status, (path, response.status_code)
        return response.json()

    def login(role):
        account = fixture["users"][role]
        client = httpx.Client(
            base_url=BASE,
            headers={
                "X-Tenant-Id": fixture["other_tenant_id"]
                if role == "outsider"
                else fixture["tenant_id"]
            },
            timeout=300,
            trust_env=False,
        )
        clients.append(client)
        session = call(
            client,
            "/api/auth/login",
            {"email": account["email"], "password": account["password"]},
        )
        client.headers["Authorization"] = "Bearer " + session["session"]["token"]
        return client

    def upload(client, thread, name, text):
        response = client.post(
            "/api/threads/" + thread + "/uploads",
            files={"files": (name, text.encode(), "text/plain")},
        )
        assert response.status_code == 200
        return response.json()["files"][0]

    def run(client, prompt, attachment, label):
        graph = "/api/langgraph/threads/" + result["thread_id"]
        attachment_ref = {"attachment_id": attachment["attachment_id"]}
        payload = {
            "assistant_id": "lead_agent",
            "input": {
                "messages": [
                    {
                        "type": "human",
                        "content": prompt,
                        "additional_kwargs": {"attachments": [attachment_ref]},
                    }
                ]
            },
            "context": {
                "agent_name": result["agent_name"],
                "subagent_enabled": False,
                "thinking_enabled": False,
                "is_plan_mode": False,
                "referenced_attachments": [attachment_ref],
            },
            "config": {"recursion_limit": 1000},
            "stream_mode": ["values", "updates"],
        }
        result["new_model_task_submissions"] += 1
        save()
        with client.stream("POST", graph + "/runs/stream", json=payload) as response:
            assert response.status_code == 200
            transcript = "\n".join(response.iter_lines())
        raw = private / (label + "-stream.txt")
        raw.write_text(transcript)
        raw.chmod(0o600)
        prior_calls = {c["id"]: c for c in result.get("actual_tool_calls", [])}
        prior_calls.update({c["id"]: c for c in tool_trace(transcript)})
        result["actual_tool_calls"] = list(prior_calls.values())
        state = call(client, graph + "/state")
        raw = private / (label + "-state.json")
        raw.write_text(json.dumps(state, ensure_ascii=False))
        raw.chmod(0o600)
        runs = call(client, graph + "/runs")
        assert runs[0]["status"] == "success" and not state.get("next"), (
            "Real graph must finish"
        )
        return state, runs[0]["run_id"]

    def verify(client, thread, paths, expected, marker, require_notes=True):
        facts = {}
        for kind, path in paths.items():
            url = "/api/threads/" + thread + "/artifacts" + path + "?download=true"
            response = client.get(url)
            assert response.status_code == 200, (kind, response.status_code)
            data = response.content
            assert "no-store" in response.headers["cache-control"]
            if kind == "xlsx":
                book = load_workbook(io.BytesIO(data), data_only=False)
                sheet = book["周收入"]
                assert [sheet.cell(row, 2).value for row in range(2, 5)] == expected
                assert marker in str(list(sheet.values))
                assert sheet._charts and type(sheet._charts[0]).__name__ == "BarChart"
                assert sheet._charts[0].series[0].val.numRef.f.endswith("$B$2:$B$4")
                book.close()
            else:
                deck = Presentation(io.BytesIO(data))
                charts = [
                    shape.chart
                    for slide in deck.slides
                    for shape in slide.shapes
                    if shape.has_chart
                ]
                assert len(charts) == 1 and list(charts[0].series[0].values) == expected
                texts = "\n".join(
                    shape.text
                    for slide in deck.slides
                    for shape in slide.shapes
                    if shape.has_text_frame
                )
                assert marker in texts and "周收入" in texts
                embedded = load_workbook(
                    io.BytesIO(charts[0].part.chart_workbook.xlsx_part.blob),
                    data_only=False,
                )
                assert [
                    embedded.active.cell(row, 2).value for row in range(2, 5)
                ] == expected
                embedded.close()
                assert not require_notes or any(
                    slide.has_notes_slide
                    and slide.notes_slide.notes_text_frame.text.strip()
                    for slide in deck.slides
                )
            facts[kind] = {
                "artifact_url": url,
                "sha256": hashlib.sha256(data).hexdigest(),
                "bytes": len(data),
                "values": expected,
                "editable_chart_verified": True,
            }
        return facts

    try:
        requester, owner, outsider = [
            login(role) for role in ["requester", "approver", "outsider"]
        ]
        if previous is None:
            agent = call(
                requester,
                "/api/agents",
                {
                    "name": "office-charts-" + suffix,
                    "template_id": "expert:office-automation",
                    "model": fixture["model_name"],
                    "tool_groups": [],
                    "skills": [],
                },
                status=201,
            )
            assert agent["presentation"]["template_id"] == "expert:office-automation"
            thread = call(requester, "/api/langgraph/threads", {})["thread_id"]
            marker = "CHART_TASK_" + secrets.token_hex(12)
            source = upload(
                requester,
                thread,
                "revenue.csv",
                "周次,收入,验收标识\n第1周,100,"
                + marker
                + "\n第2周,150,\n第3周,120,\n",
            )
            result.update(
                agent_id=agent["id"],
                agent_name=agent["name"],
                thread_id=thread,
                marker=marker,
                template_id=agent["presentation"]["template_id"],
                template_prompt_sha256=hashlib.sha256(
                    agent["system_prompt"].encode()
                ).hexdigest(),
                source_attachment=source,
                output_prefix="/mnt/user-data/outputs/charts-" + suffix,
            )
            save()
        thread, marker, prefix = (
            result["thread_id"],
            result["marker"],
            result["output_prefix"],
        )
        initial = result.get("initial_paths") or {
            kind: prefix + "-initial." + kind for kind in ["xlsx", "pptx"]
        }
        revised = {kind: prefix + "-revised." + kind for kind in ["xlsx", "pptx"]}
        if previous and "initial_task_completed" not in result["phases"]:
            # Preserve successful real chart data. Repair only the missing notes
            # in the same Agent/thread; do not regenerate the initial files.
            result["original_initial_artifacts"] = verify(
                requester, thread, initial, [100, 150, 120], marker, require_notes=False
            )
            original_pptx = initial["pptx"]
            initial = {**initial, "pptx": prefix + "-initial-notes.pptx"}
            prompt = (
                "请检查本线程已经生成的 "
                + original_pptx
                + "，为第2页新增演讲者备注：说明收入数据来自revenue.csv附件且未经外部验证。只修改演讲者备注，不要重建文档或修改图表、正文与数据。使用原生Office工具另存为 "
                + initial["pptx"]
                + "，检查备注并提供下载链接。"
            )
            state, run_id = run(
                requester, prompt, result["source_attachment"], "notes-recovery"
            )
            result["initial_artifacts"] = verify(
                requester, thread, initial, [100, 150, 120], marker
            )
            result.update(
                initial_paths=initial,
                initial_run_id=run_id,
                initial_message_ids=[m.get("id") for m in state["values"]["messages"]],
            )
            result["phases"].extend(
                [
                    "notes_recovered_without_regenerating_charts",
                    "initial_task_completed",
                ]
            )
            save()
        if "initial_task_completed" not in result["phases"]:
            source = result["source_attachment"]
            prompt = (
                "请读取附件 "
                + source["virtual_path"]
                + "，根据三周真实收入制作可编辑图表汇报包。Excel工作表名为“周收入”，A列周次、B列数值收入，保留附件验收标识，增加收入柱状图。PowerPoint至少两页：标题页含“周收入”和原样验收标识，图表页含同样数据的可编辑柱状图，演讲者备注中说明来源。分别保存为 "
                + initial["xlsx"]
                + " 和 "
                + initial["pptx"]
                + "，提供下载链接；不编造数据，不用图片替代图表。"
            )
            state, run_id = run(requester, prompt, source, "initial")
            result["initial_artifacts"] = verify(
                requester, thread, initial, [100, 150, 120], marker
            )
            result["initial_paths"] = initial
            result["initial_run_id"] = run_id
            result["initial_message_ids"] = [
                m.get("id") for m in state["values"]["messages"]
            ]
            result["phases"].append("initial_task_completed")
            save()
        if "revision_task_completed" not in result["phases"]:
            correction = upload(
                requester,
                thread,
                "correction.txt",
                "收入更正：第2周由150改为175；第1周100和第3周120保持。保留原始文件与验收标识。",
            )
            prompt = (
                "请读取更正附件 "
                + correction["virtual_path"]
                + "，检查刚才两份文件后，修改 Excel "
                + initial["xlsx"]
                + " 的原数据并保留原柱状图；修改 PowerPoint "
                + initial["pptx"]
                + " 的原可编辑柱状图及内嵌工作表，保留其他页面和备注。不要重建替代原文档。结果另存为 "
                + revised["xlsx"]
                + " 和 "
                + revised["pptx"]
                + "，不要覆盖原文件；检查结果并提供两份下载链接。"
            )
            state, run_id = run(requester, prompt, correction, "revision")
            result["revised_artifacts"] = verify(
                requester, thread, revised, [100, 175, 120], marker
            )
            calls = result["actual_tool_calls"]
            assert any(
                c["name"] == "office_xlsx" and c.get("action") == "edit" for c in calls
            )
            assert any(
                c["name"] == "office_pptx" and c.get("action") == "edit_charts"
                for c in calls
            )
            assert {"read_file", "present_files", "office_xlsx", "office_pptx"} <= {
                c["name"] for c in calls
            }
            result.update(revision_run_id=run_id, tool_names=[c["name"] for c in calls])
            result["phases"].append("revision_task_completed")
            save()
        # Rechecks never initiate new model requests once both stages are accepted.
        assert (
            verify(requester, thread, initial, [100, 150, 120], marker)
            == result["initial_artifacts"]
        )
        assert (
            verify(requester, thread, revised, [100, 175, 120], marker)
            == result["revised_artifacts"]
        )
        for facts in [result["initial_artifacts"], result["revised_artifacts"]]:
            for artifact in facts.values():
                assert owner.get(artifact["artifact_url"]).status_code == 404
                assert outsider.get(artifact["artifact_url"]).status_code == 404
        graph = "/api/langgraph/threads/" + thread
        for other in [owner, outsider]:
            assert other.get(graph + "/state").status_code == 404
        requester.close()
        requester = login("requester")
        history = call(requester, graph + "/history", {"limit": 30})
        if result.get("first_task_checkpoint"):
            first = call(
                requester,
                graph + "/state/checkpoint",
                {"checkpoint": result["first_task_checkpoint"]},
            )
            assert [m.get("id") for m in first["values"]["messages"]] == result[
                "first_task_message_ids"
            ]
            result["first_task_checkpoint_preserved"] = True
            history.append(first)
        assert all(
            path in json.dumps(history)
            for path in [*initial.values(), *revised.values()]
        )
        assert (
            verify(requester, thread, revised, [100, 175, 120], marker)
            == result["revised_artifacts"]
        )
        agent = call(requester, "/api/agents/" + result["agent_name"])
        assert (
            hashlib.sha256(agent["system_prompt"].encode()).hexdigest()
            == result["template_prompt_sha256"]
        )
        state = call(requester, graph + "/state")
        assert (
            not state.get("next")
            and call(requester, graph + "/runs")[0]["status"] == "success"
        )
        result.update(
            complete=True,
            run_status="success",
            fresh_login_history_and_artifacts=True,
            current_acl_verified=True,
            original_files_and_template_preserved=True,
        )
    except Exception as error:
        result.update(failure_type=type(error).__name__, failure=str(error))
        raise
    finally:
        for client in clients:
            client.close()
        save()
    print("Actual Office chart create/revise task accepted")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--resume", type=Path)
    args = parser.parse_args()
    execute(args.output, args.resume)
